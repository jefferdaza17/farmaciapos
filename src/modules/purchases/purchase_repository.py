"""Persistencia de compras en Excel."""

from pathlib import Path

from openpyxl import Workbook, load_workbook


class PurchaseRepository:
    """Guarda las entradas de productos al inventario."""

    SHEET_NAME = "Compras"
    HEADERS = (
        "Compra",
        "Fecha",
        "Proveedor",
        "Código",
        "Cantidad",
        "Costo unitario",
        "Total",
    )

    def __init__(self, file_path: Path | None = None) -> None:
        project_root = Path(__file__).resolve().parents[3]
        self._file_path = file_path or project_root / "data" / "excel" / "farmacia.xlsx"

    def save_purchase(
        self,
        purchase_id: str,
        date: str,
        supplier: str,
        code: str,
        quantity: int,
        unit_cost: float,
    ) -> None:
        """Registra una entrada de inventario."""
        self._ensure_sheet()
        workbook = load_workbook(self._file_path)
        workbook[self.SHEET_NAME].append(
            (
                purchase_id,
                date,
                supplier,
                code,
                quantity,
                unit_cost,
                quantity * unit_cost,
            )
        )
        workbook.save(self._file_path)
        workbook.close()

    def _ensure_sheet(self) -> None:
        self._file_path.parent.mkdir(parents=True, exist_ok=True)
        if self._file_path.exists():
            workbook = load_workbook(self._file_path)
        else:
            workbook = Workbook()
            workbook.remove(workbook.active)
        if self.SHEET_NAME not in workbook.sheetnames:
            worksheet = workbook.create_sheet(self.SHEET_NAME)
            worksheet.append(self.HEADERS)
            worksheet.freeze_panes = "A2"
        workbook.save(self._file_path)
        workbook.close()
