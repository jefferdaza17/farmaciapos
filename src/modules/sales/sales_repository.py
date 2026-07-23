"""Persistencia de ventas y sus detalles en Excel."""

from pathlib import Path

from openpyxl import Workbook, load_workbook


class SalesRepository:
    """Guarda ventas confirmadas en el mismo libro de datos de la farmacia."""

    SALES_SHEET = "Ventas"
    DETAILS_SHEET = "Detalle ventas"
    SALES_HEADERS = ("Factura", "Fecha", "Total")
    DETAIL_HEADERS = (
        "Factura",
        "Código",
        "Nombre",
        "Cantidad",
        "Precio unitario",
        "Subtotal",
    )

    def __init__(self, file_path: Path | None = None) -> None:
        project_root = Path(__file__).resolve().parents[3]
        self._file_path = file_path or project_root / "data" / "excel" / "farmacia.xlsx"

    def save_sale(
        self, receipt: str, date: str, items: list[dict[str, object]]
    ) -> None:
        """Registra la cabecera y cada producto de una venta."""
        self._ensure_sheets()
        total = sum(float(item["Subtotal"]) for item in items)
        workbook = load_workbook(self._file_path)
        workbook[self.SALES_SHEET].append((receipt, date, total))
        details = workbook[self.DETAILS_SHEET]
        for item in items:
            details.append(
                (
                    receipt,
                    item["Código"],
                    item["Nombre"],
                    item["Cantidad"],
                    item["Precio"],
                    item["Subtotal"],
                )
            )
        workbook.save(self._file_path)
        workbook.close()

    def _ensure_sheets(self) -> None:
        self._file_path.parent.mkdir(parents=True, exist_ok=True)
        if self._file_path.exists():
            workbook = load_workbook(self._file_path)
        else:
            workbook = Workbook()
            workbook.remove(workbook.active)

        for sheet_name, headers in (
            (self.SALES_SHEET, self.SALES_HEADERS),
            (self.DETAILS_SHEET, self.DETAIL_HEADERS),
        ):
            if sheet_name not in workbook.sheetnames:
                worksheet = workbook.create_sheet(sheet_name)
                worksheet.append(headers)
                worksheet.freeze_panes = "A2"
        workbook.save(self._file_path)
        workbook.close()
