"""Persistencia de clientes y proveedores en Excel."""

from pathlib import Path

from openpyxl import Workbook, load_workbook


class ContactRepository:
    """Administra una hoja de contactos dentro del libro de la farmacia."""

    HEADERS = ("Documento", "Nombre", "Teléfono", "Correo")

    def __init__(self, sheet_name: str, file_path: Path | None = None) -> None:
        project_root = Path(__file__).resolve().parents[3]
        self._sheet_name = sheet_name
        self._file_path = file_path or project_root / "data" / "excel" / "farmacia.xlsx"

    def load(self) -> list[dict[str, object]]:
        """Devuelve los contactos guardados."""
        self._ensure_sheet()
        workbook = load_workbook(self._file_path, data_only=True)
        worksheet = workbook[self._sheet_name]
        rows = list(worksheet.iter_rows(min_row=2, values_only=True))
        workbook.close()
        return [
            dict(zip(self.HEADERS, row, strict=True))
            for row in rows
            if any(value is not None for value in row)
        ]

    def save(self, contacts: list[dict[str, object]]) -> None:
        """Guarda el listado completo de contactos."""
        self._ensure_sheet()
        workbook = load_workbook(self._file_path)
        worksheet = workbook[self._sheet_name]
        worksheet.delete_rows(2, worksheet.max_row)
        for contact in contacts:
            worksheet.append([contact[header] for header in self.HEADERS])
        workbook.save(self._file_path)
        workbook.close()

    def _ensure_sheet(self) -> None:
        self._file_path.parent.mkdir(parents=True, exist_ok=True)
        if self._file_path.exists():
            workbook = load_workbook(self._file_path)
        else:
            workbook = Workbook()
            workbook.remove(workbook.active)
        if self._sheet_name not in workbook.sheetnames:
            worksheet = workbook.create_sheet(self._sheet_name)
            worksheet.append(self.HEADERS)
            worksheet.freeze_panes = "A2"
        workbook.save(self._file_path)
        workbook.close()
