"""Consultas agregadas para dashboard y reportes."""

from pathlib import Path

from openpyxl import load_workbook

from src.modules.inventory.inventory_service import InventoryService


class ReportService:
    """Calcula indicadores a partir del inventario y el libro Excel."""

    def __init__(
        self,
        inventory_service: InventoryService | None = None,
        file_path: Path | None = None,
    ) -> None:
        self._inventory_service = inventory_service or InventoryService()
        project_root = Path(__file__).resolve().parents[3]
        self._file_path = file_path or project_root / "data" / "excel" / "farmacia.xlsx"

    def summary(self) -> dict[str, float | int]:
        """Devuelve los indicadores principales del negocio."""
        medications = self._inventory_service.list_medications()
        return {
            "Medicamentos registrados": len(medications),
            "Unidades en inventario": sum(int(item["Stock"]) for item in medications),
            "Productos con stock bajo": sum(
                int(item["Stock"]) <= int(item["Stock mínimo"]) for item in medications
            ),
            "Ventas registradas": self._row_count("Ventas"),
            "Total vendido": self._sum_column("Ventas", 3),
            "Compras registradas": self._row_count("Compras"),
            "Total comprado": self._sum_column("Compras", 7),
        }

    def low_stock(self) -> list[dict[str, object]]:
        """Lista los medicamentos que requieren reposición."""
        return [
            item
            for item in self._inventory_service.list_medications()
            if int(item["Stock"]) <= int(item["Stock mínimo"])
        ]

    def _row_count(self, sheet_name: str) -> int:
        if not self._file_path.exists():
            return 0
        workbook = load_workbook(self._file_path, read_only=True, data_only=True)
        if sheet_name not in workbook.sheetnames:
            workbook.close()
            return 0
        count = max(0, workbook[sheet_name].max_row - 1)
        workbook.close()
        return count

    def _sum_column(self, sheet_name: str, column: int) -> float:
        if not self._file_path.exists():
            return 0.0
        workbook = load_workbook(self._file_path, read_only=True, data_only=True)
        if sheet_name not in workbook.sheetnames:
            workbook.close()
            return 0.0
        total = sum(
            float(row[0] or 0)
            for row in workbook[sheet_name].iter_rows(
                min_row=2, min_col=column, max_col=column, values_only=True
            )
        )
        workbook.close()
        return total
