"""Persistencia del inventario en el archivo Excel local."""

from copy import copy
from pathlib import Path

from openpyxl import Workbook, load_workbook


class InventoryRepository:
    """Lee y guarda medicamentos en ``data/excel/farmacia.xlsx``."""

    SHEET_NAME = "Inventario"
    HEADERS = (
        "Código",
        "Nombre",
        "Laboratorio",
        "Precio",
        "Stock",
        "Stock mínimo",
    )

    def __init__(self, file_path: Path | None = None) -> None:
        project_root = Path(__file__).resolve().parents[3]
        self._file_path = file_path or project_root / "data" / "excel" / "farmacia.xlsx"

    def load(self) -> list[dict[str, object]]:
        """Devuelve todos los medicamentos almacenados."""
        self._ensure_file()
        workbook = load_workbook(self._file_path, data_only=True)
        worksheet = workbook[self.SHEET_NAME]
        rows = list(worksheet.iter_rows(min_row=2, values_only=True))
        workbook.close()

        return [
            dict(zip(self.HEADERS, row, strict=True))
            for row in rows
            if any(value is not None for value in row)
        ]

    def save(self, medications: list[dict[str, object]]) -> None:
        """Reemplaza el contenido del inventario con los datos recibidos."""
        self._ensure_file()
        workbook = load_workbook(self._file_path)
        worksheet = workbook[self.SHEET_NAME]
        worksheet.delete_rows(2, worksheet.max_row)
        for medication in medications:
            worksheet.append([medication[header] for header in self.HEADERS])
        workbook.save(self._file_path)
        workbook.close()

    def _ensure_file(self) -> None:
        """Crea el libro y su hoja de inventario cuando aún no existen."""
        if self._file_path.exists():
            return

        self._file_path.parent.mkdir(parents=True, exist_ok=True)
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = self.SHEET_NAME
        worksheet.append(self.HEADERS)
        worksheet.freeze_panes = "A2"
        for cell in worksheet[1]:
            font = copy(cell.font)
            font.bold = True
            cell.font = font
        workbook.save(self._file_path)
        workbook.close()
